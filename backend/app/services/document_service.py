
import csv
import io
from pathlib import Path

import pymupdf
from docx import Document


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}


def extract_text(file_path: str) -> list[dict]:
    """Extract text and return sections with page or row information."""

    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type.")

    sections = []

    if extension == ".pdf":
        with pymupdf.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()

                if text:
                    sections.append({
                        "location": f"Page {page_number}",
                        "text": text,
                    })

    elif extension == ".docx":
        doc = Document(path)
        content = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                content.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                content.append(" | ".join(cells))

        if content:
            sections.append({
                "location": "Document body",
                "text": "\n".join(content),
            })

    elif extension == ".txt":
        text = path.read_text(encoding="utf-8-sig").strip()

        if text:
            sections.append({
                "location": "Text file",
                "text": text,
            })

    elif extension == ".csv":
        with path.open(
            "r", encoding="utf-8-sig", newline=""
        ) as file:
            reader = csv.reader(file)
            headers = next(reader, None)

            for row_number, row in enumerate(reader, start=2):
                if headers:
                    content = [
                        f"{header}: {value}"
                        for header, value in zip(headers, row)
                    ]
                else:
                    content = row

                if content:
                    sections.append({
                        "location": f"Row {row_number}",
                        "text": " | ".join(content),
                    })

    return sections
